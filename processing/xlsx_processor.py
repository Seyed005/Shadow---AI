from pathlib import Path

from openpyxl import load_workbook


def extract_xlsx_content(xlsx_path: str) -> dict:
    """
    Extract text from an XLSX workbook while preserving
    worksheet and cell-level source provenance.

    The returned canonical text is clean and suitable
    for the existing Shadow AI detection pipeline.

    Each source segment records the exact character range
    belonging to an original worksheet cell.
    """

    xlsx_file = Path(xlsx_path)

    if not xlsx_file.exists():
        raise FileNotFoundError(
            f"XLSX file not found: {xlsx_file}"
        )

    if xlsx_file.suffix.lower() != ".xlsx":
        raise ValueError(
            "Input file must be an XLSX."
        )

    workbook = load_workbook(
        filename=xlsx_path,
        data_only=False
    )

    worksheets = []
    source_segments = []

    canonical_parts = []
    current_position = 0

    def add_cell_segment(
        text: str,
        sheet_name: str,
        cell_coordinate: str
    ):
        nonlocal current_position

        if not text:
            return

        if canonical_parts:
            canonical_parts.append("\n")
            current_position += 1

        start = current_position

        canonical_parts.append(text)
        current_position += len(text)

        end = current_position

        source_segments.append({
            "start": start,
            "end": end,
            "location": (
                f"SHEET '{sheet_name}' "
                f"CELL {cell_coordinate}"
            ),
            "source_type": "XLSX_CELL",
            "sheet_name": sheet_name,
            "cell": cell_coordinate
        })

    # ---------------------------------------------------------
    # WORKSHEETS
    # ---------------------------------------------------------

    for worksheet in workbook.worksheets:

        sheet_name = worksheet.title

        sheet_data = {
            "sheet_name": sheet_name,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "cells": []
        }

        for row in worksheet.iter_rows():

            for cell in row:

                if cell.value is None:
                    continue

                cell_text = str(cell.value).strip()

                if not cell_text:
                    continue

                sheet_data["cells"].append({
                    "cell": cell.coordinate,
                    "value": cell_text
                })

                add_cell_segment(
                    text=cell_text,
                    sheet_name=sheet_name,
                    cell_coordinate=cell.coordinate
                )

        worksheets.append(sheet_data)

    canonical_text = "".join(
        canonical_parts
    ).strip()

    workbook.close()

    return {
        "file_name": xlsx_file.name,
        "worksheet_count": len(worksheets),
        "worksheets": worksheets,
        "text": canonical_text,
        "source_segments": source_segments
    }


if __name__ == "__main__":

    test_xlsx = (
        "sample_data/"
        "synthetic_attacks/"
        "test_sensitive.xlsx"
    )

    try:

        result = extract_xlsx_content(
            test_xlsx
        )

        print("\nXLSX PROCESSING RESULT")
        print("======================")

        print(
            "File:",
            result["file_name"]
        )

        print(
            "Worksheets:",
            result["worksheet_count"]
        )

        print("\nCanonical Text:")
        print(result["text"])

        print("\nSource Segments:")

        for segment in result[
            "source_segments"
        ]:
            print(segment)

    except Exception as error:

        print("\nXLSX PROCESSING ERROR")
        print("====================")
        print(error)