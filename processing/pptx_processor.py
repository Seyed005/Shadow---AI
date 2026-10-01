from pathlib import Path

from pptx import Presentation


def extract_pptx_content(pptx_path: str) -> dict:
    """
    Extract text from a PPTX presentation while preserving
    slide and shape-level source provenance.

    The returned canonical text is clean and suitable
    for the existing Shadow AI detection pipeline.

    Each source segment records the exact character range
    belonging to an original slide/shape.
    """

    pptx_file = Path(pptx_path)

    if not pptx_file.exists():
        raise FileNotFoundError(
            f"PPTX file not found: {pptx_file}"
        )

    if pptx_file.suffix.lower() != ".pptx":
        raise ValueError(
            "Input file must be a PPTX."
        )

    presentation = Presentation(pptx_path)

    slides = []
    source_segments = []

    canonical_parts = []
    current_position = 0

    def add_shape_segment(
        text: str,
        slide_number: int,
        shape_number: int,
        shape_type: str
    ):
        nonlocal current_position

        if not text:
            return

        if canonical_parts:
            canonical_parts.append("\n")
            current_position += 1

        start = current_position

        canonical_parts.append(text)
        current_position += len(text)

        end = current_position

        source_segments.append({
            "start": start,
            "end": end,
            "location": (
                f"SLIDE {slide_number} "
                f"SHAPE {shape_number}"
            ),
            "source_type": "PPTX_SHAPE",
            "slide_number": slide_number,
            "shape_number": shape_number,
            "shape_type": shape_type
        })

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_data = {
            "slide_number": slide_number,
            "shapes": []
        }

        for shape_number, shape in enumerate(
            slide.shapes,
            start=1
        ):

            shape_text = ""

            # Normal text boxes/placeholders
            if hasattr(shape, "text"):
                shape_text = shape.text.strip()

            # Tables may not expose all content through
            # shape.text, so extract each table cell.
            if getattr(shape, "has_table", False):

                table_rows = []

                for row_number, row in enumerate(
                    shape.table.rows,
                    start=1
                ):

                    row_cells = []

                    for cell_number, cell in enumerate(
                        row.cells,
                        start=1
                    ):

                        cell_text = cell.text.strip()

                        row_cells.append({
                            "cell_number": cell_number,
                            "text": cell_text
                        })

                    table_rows.append({
                        "row_number": row_number,
                        "cells": row_cells
                    })

                table_text_parts = []

                for row in table_rows:
                    values = [
                        cell["text"]
                        for cell in row["cells"]
                    ]

                    row_text = " | ".join(values)

                    if row_text:
                        table_text_parts.append(row_text)

                shape_text = "\n".join(
                    table_text_parts
                ).strip()

            if not shape_text:
                continue

            shape_type = str(
                getattr(
                    shape,
                    "shape_type",
                    "UNKNOWN"
                )
            )

            slide_data["shapes"].append({
                "shape_number": shape_number,
                "shape_type": shape_type,
                "text": shape_text
            })

            add_shape_segment(
                text=shape_text,
                slide_number=slide_number,
                shape_number=shape_number,
                shape_type=shape_type
            )

        slides.append(slide_data)

    canonical_text = "".join(
        canonical_parts
    ).strip()

    return {
        "file_name": pptx_file.name,
        "slide_count": len(slides),
        "slides": slides,
        "text": canonical_text,
        "source_segments": source_segments
    }


if __name__ == "__main__":

    test_pptx = (
        "sample_data/"
        "synthetic_attacks/"
        "test_sensitive.pptx"
    )

    try:

        result = extract_pptx_content(
            test_pptx
        )

        print("\nPPTX PROCESSING RESULT")
        print("======================")

        print(
            "File:",
            result["file_name"]
        )

        print(
            "Slides:",
            result["slide_count"]
        )

        print("\nCanonical Text:")
        print(result["text"])

        print("\nSource Segments:")

        for segment in result[
            "source_segments"
        ]:
            print(segment)

    except Exception as error:

        print("\nPPTX PROCESSING ERROR")
        print("====================")
        print(error)