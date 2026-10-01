from pathlib import Path

from docx import Document


def extract_docx_content(docx_path: str) -> dict:
    """
    Extract text from a DOCX document and preserve
    source-location ranges for provenance mapping.

    The canonical text remains clean (no artificial
    paragraph/table markers), while source_segments
    records exactly where each piece of text came from.
    """

    docx_file = Path(docx_path)

    if not docx_file.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {docx_file}"
        )

    if docx_file.suffix.lower() != ".docx":
        raise ValueError(
            "Input file must be a DOCX."
        )

    document = Document(docx_path)

    paragraphs = []
    tables = []
    source_segments = []

    canonical_parts = []
    current_position = 0

    def add_text_segment(
        text: str,
        location: str,
        source_type: str,
        source_id: str
    ):
        nonlocal current_position

        if not text:
            return

        # Add a newline between logical source sections.
        if canonical_parts:
            canonical_parts.append("\n")
            current_position += 1

        start = current_position

        canonical_parts.append(text)
        current_position += len(text)

        end = current_position

        source_segments.append({
            "start": start,
            "end": end,
            "location": location,
            "source_type": source_type,
            "source_id": source_id
        })

    # ---------------------------------------------------------
    # PARAGRAPHS
    # ---------------------------------------------------------

    for index, paragraph in enumerate(
        document.paragraphs,
        start=1
    ):
        text = paragraph.text.strip()

        if not text:
            continue

        paragraphs.append({
            "paragraph_number": index,
            "text": text
        })

        add_text_segment(
            text=text,
            location=f"PARAGRAPH {index}",
            source_type="PARAGRAPH",
            source_id=str(index)
        )

    # ---------------------------------------------------------
    # TABLES
    # ---------------------------------------------------------

    for table_number, table in enumerate(
        document.tables,
        start=1
    ):
        rows = []

        for row_number, row in enumerate(
            table.rows,
            start=1
        ):
            cells = []

            for cell_number, cell in enumerate(
                row.cells,
                start=1
            ):
                cell_text = cell.text.strip()

                cells.append({
                    "cell_number": cell_number,
                    "text": cell_text
                })

            rows.append({
                "row_number": row_number,
                "cells": cells
            })

            # Keep the table row in the canonical text
            # as a single searchable representation.
            row_values = [
                cell["text"]
                for cell in cells
            ]

            row_text = " | ".join(row_values)

            if row_text:
                add_text_segment(
                    text=row_text,
                    location=(
                        f"TABLE {table_number} "
                        f"ROW {row_number}"
                    ),
                    source_type="TABLE_ROW",
                    source_id=(
                        f"{table_number}:{row_number}"
                    )
                )

        tables.append({
            "table_number": table_number,
            "rows": rows
        })

    canonical_text = "".join(
        canonical_parts
    ).strip()

    return {
        "file_name": docx_file.name,
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "paragraphs": paragraphs,
        "tables": tables,
        "text": canonical_text,
        "source_segments": source_segments
    }


if __name__ == "__main__":

    test_docx = (
        "sample_data/"
        "synthetic_attacks/"
        "test_sensitive.docx"
    )

    try:
        result = extract_docx_content(test_docx)

        print("\nDOCX PROCESSING RESULT")
        print("======================")

        print(
            "File:",
            result["file_name"]
        )

        print(
            "Paragraphs:",
            result["paragraph_count"]
        )

        print(
            "Tables:",
            result["table_count"]
        )

        print("\nCanonical Text:")
        print(result["text"])

        print("\nSource Segments:")

        for segment in result["source_segments"]:
            print(segment)

    except Exception as error:

        print("\nDOCX PROCESSING ERROR")
        print("====================")
        print(error)